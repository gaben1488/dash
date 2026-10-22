"""An answer is evidence of what was answered, never a made-up current result."""
import hashlib
from dataclasses import asdict, replace
from io import BytesIO

from docx import Document
from test_identity_and_metrics import row


def source_record(decision='Закупка удалена из плана-реестра'):
    from procurement_engine.recommendation_links import verify_saved_report_origin

    text = 'Вынести на ЭА 1 Поставка бумаги – 12,00 тыс. руб.'
    doc = Document(); doc.add_paragraph('ОТЧЕТ'); doc.add_paragraph('срез на 25.09.2026'); doc.add_paragraph('УО')
    table = doc.add_table(rows=2, cols=4)
    for cell, value in zip(table.rows[0].cells, ['№ п/п', 'Рекомендация', 'Ответ ГРБС на рекомендацию', 'Решение УЭР']):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ['1', text, 'Позиция принята УО', decision + '\n25.09: УДАЛЕНО']):
        cell.text = value
    stream = BytesIO(); doc.save(stream); content = stream.getvalue(); digest = hashlib.sha256(content).hexdigest()
    rec = {'grbs': 'УО', 'recommendation_text': text, 'grbs_response_original': 'Позиция принята УО',
        'uer_decision_original': decision, 'origin_evidence': [{'kind': 'SAVED_REPORT_RECOMMENDATION_TEXT',
            'document_sha256': digest, 'document_date': '2026-09-25', 'table': 1, 'row': 2, 'cell': 2,
            'grbs_heading': 'УО', 'text': text, 'text_sha256': hashlib.sha256(text.encode()).hexdigest()}]}
    docs = {digest: content}
    return rec, verify_saved_report_origin(rec, docs), docs


def test_dated_answer_is_verified_separately_and_does_not_become_current_execution():
    from procurement_engine.answered_questions import saved_answer

    record, proof, docs = source_record()
    answer = saved_answer(record, proof, docs)
    assert answer['kind'] == 'WITHDRAWN_REPORTED'
    assert answer['as_of'] == '2026-09-25'
    assert answer['proves_current_implementation'] is False
    assert '25.09:' not in answer['decision']
    assert saved_answer({**record, 'uer_decision_original': 'Заключено с ЕП'}, proof, docs) is None
    assert saved_answer(record, proof, {}) is None
    assert saved_answer(record, None, docs) is None


def test_plain_acceptance_is_not_an_answer_to_implementation():
    from procurement_engine.answered_questions import saved_answer

    record, proof, docs = source_record('Позиция принята УО')
    assert saved_answer(record, proof, docs) is None


def test_group_answer_requires_every_member_and_preserves_current_supplement():
    from procurement_engine.answered_questions import _covered_group_answer

    record = {'grbs': 'УО', 'recommendation_text': 'Вынести на ЭА 1,2 бензин, ГСМ – 200,00 тыс. руб.'}
    answer = '1 заключено с еп'
    assert not _covered_group_answer(record, answer, [])
    current = row(source_row_no='2', subject='Бензин', procurement_uid='PUR-two', grbs_comment='договор заключен')
    result = _covered_group_answer(record, answer, [current])
    assert result['current_source_statements'][0]['quote'] == 'договор заключен'
    assert not _covered_group_answer(record, answer, [replace(current, subject='Другой бензин')])
    assert not _covered_group_answer(record, answer, [replace(current, procurement_uid=None)])
    assert not _covered_group_answer(record, 'если ' + answer, [current])
    assert not _covered_group_answer(record, 'не подтверждено, что ' + answer, [current])


def test_reviewed_joint_customer_alias_is_scoped_and_money_suffix_is_not_a_subject_word():
    from procurement_engine.recommendation_links import (
        _registered_joint_customer,
        literal_plan_label,
    )

    assert _registered_joint_customer('УО', 'Совместные закупки (УО)')
    assert not _registered_joint_customer('УО', 'Совместные закупки (УКСиМП)')
    assert literal_plan_label('Замена ворот (из 405,00 - 121,00 это Собственные деньги с платных услуг)') == 'Замена ворот'
    assert literal_plan_label('Замена ворот (другой гараж)') == 'Замена ворот (другой гараж)'


def test_literal_label_keeps_program_prefix_and_full_current_name_without_substring_guess():
    from procurement_engine.answered_questions import literal_action
    from procurement_engine.recommendation_links import literal_plan_label

    subject = '2.2.3.1 пледы / поставка пледов в подарочных коробках'
    r = row(subject=subject)
    assert literal_plan_label(subject) == 'пледы'
    record = {'recommendation_text': 'Вынести на ЭА 1 пледы – 164,29 тыс. руб'}
    assert literal_action(record, [r], ['1'])['type'] == 'CHANGE_METHOD_EA'
    assert literal_action({'recommendation_text': record['recommendation_text'] + ' если будет финансирование'}, [r], ['1']) is None
    assert literal_action(record, [replace(r, subject='Пледы для других целей')], ['1']) is None


def test_joint_scope_discussion_must_quote_the_exact_expanded_subject_without_conflicting_replies():
    from procurement_engine.answered_questions import saved_discussion_answer

    record = {'grbs': 'УО', 'recommendation_type': 'MERGE_PROCUREMENTS',
        'recommendation_text': 'Объединить «Оказание услуг по изготовлению баннерной продукции в рамках ярмарки»'}
    r = row(subject='Оказание услуг по изготовлению и монтажу/демонтажу баннерной продукции в рамках ярмарки',
            procurement_uid='PUR-test')
    discussion = {'id': 'thread', 'content': 'Объединение с мероприятием по монтажу/демонтажу баннерной продукции',
                  'quotedFileContent': {'value': r.subject}}
    source = {'grbs': 'УО', 'role': 'master', 'provider_id': r.source_id, 'sheet': r.sheet_name,
              'formula_evidence': {'discussions': [discussion]}}
    origin = {'document_sha256': 'a' * 64, 'table': 1, 'row': 2}
    answer = saved_discussion_answer(record, origin, [r], [source], '21.10.2026')
    assert answer['kind'] == 'GROUP_SCOPE_REPORTED'
    assert answer['discussion_evidence']['id'] == 'thread'
    assert answer['proves_current_implementation'] is False
    assert saved_discussion_answer(record, origin, [replace(r, subject='Другая закупка')], [source], '21.10.2026') is None
    discussion['replies'] = [{'content': 'Объединение отменено'}]
    assert saved_discussion_answer(record, origin, [r], [source], '21.10.2026') is None
    assert saved_discussion_answer(record, None, [r], [source], '21.10.2026') is None


def test_existing_yearlong_stage_is_neutral_without_inventing_kind_or_completion():
    from procurement_engine.automation_assurance import assess_automation

    r = row(method='ЕП', actual_date=None, fact_mb=12, procurement_uid='PUR-test')
    detail = {**asdict(r), 'physical_row_key': r.physical_row_key}
    values = [[], [], [], [''] * 17]
    values[3][16] = 'х'
    model = {'details': [detail], 'recommendation_records': [], 'issues': []}
    source = [{'provider_id': r.source_id, 'sheet': r.sheet_name, 'values': values}]
    result = assess_automation(model, source, contract='actionable-assurance-v4')
    assert result['user_action_count'] == 0
    assert result['resolved_questions'][0]['code'] == 'YEARLONG_STAGE'
    assert result['resolved_questions'][0]['proves_current_implementation'] is False
    assert detail['actual_date'] is None
    assert assess_automation(model, source, contract='actionable-assurance-v3')['user_action_count'] == 1
    values[3][16] = 'неизвестно'
    assert assess_automation(model, source, contract='actionable-assurance-v4')['user_action_count'] == 1
