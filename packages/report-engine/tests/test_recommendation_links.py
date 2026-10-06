import hashlib
from dataclasses import replace

import pytest
from procurement_engine.models import ProcurementRow
from procurement_engine.recommendation_links import resolve_current_link

TEXT = 'Рекомендуем позицию 42 вынести на ЭА (Поставка бумаги) на сумму 46,00 тыс. руб.'
TEST_DOCUMENTS = {}


def recommendation(text=TEXT):
    from io import BytesIO

    from docx import Document

    doc = Document(); doc.add_paragraph('ОТЧЕТ ПО ЗАКУПКАМ'); doc.add_paragraph('срез на 25.09.2026')
    doc.add_paragraph('УЭР'); doc.add_table(rows=2, cols=2).cell(1, 1).text = text
    buffer = BytesIO(); doc.save(buffer); content = buffer.getvalue(); digest = hashlib.sha256(content).hexdigest()
    TEST_DOCUMENTS[digest] = content
    return {'recommendation_id': 'synthetic', 'grbs': 'УЭР', 'recommendation_text': text,
            'source_procurement_ids': ['999'], 'semantic_status': 'IMPLEMENTED',
            'origin_evidence': [{'kind': 'SAVED_REPORT_RECOMMENDATION_TEXT',
                'document_sha256': digest, 'document_date': '2026-09-25',
                'table': 1, 'row': 2, 'cell': 2, 'grbs_heading': 'УЭР',
                'text': text, 'text_sha256': hashlib.sha256(text.encode()).hexdigest()}]}


def row(**kw):
    return ProcurementRow(snapshot_id='snapshot', procurement_id='42', source_row_no='42',
        source_id='synthetic-book', sheet_name='ALL', row_number=4, grbs='УЭР',
        subject='Поставка бумаги', method='ЭА', plan_mb=46, procurement_uid='PUR-synthetic', **kw)


def resolve(rec, rows):
    from procurement_engine.recommendation_links import verify_saved_report_origin

    proof = verify_saved_report_origin(rec, TEST_DOCUMENTS)
    return resolve_current_link(rec, rows, report_date='30.09.2026', snapshot_id='snapshot', verified_origin=proof)


def test_explicit_text_subject_money_and_persisted_uid_establish_current_link_only():
    result = resolve(recommendation(), [row()])
    assert result['status'] == 'CONFIRMED'
    assert result['procurement_uids'] == ['PUR-synthetic']
    assert result['business_ids'] == ['42']
    assert result['fulfillment'] == 'UNKNOWN'
    assert result['evidence_snapshot_id'] == 'snapshot'


@pytest.mark.parametrize('change', [{'subject': 'Другая закупка'}, {'plan_mb': 47},
    {'grbs': 'УО'}, {'planned_year': 2027}, {'procurement_uid': None}])
def test_reused_number_changed_scope_or_unresolved_identity_is_not_a_link(change):
    assert resolve(recommendation(), [replace(row(), **change)])['status'] != 'CONFIRMED'


def test_multiple_equal_current_candidates_remain_ambiguous():
    assert resolve(recommendation(), [row(), replace(row(), row_number=5, procurement_uid='PUR-other')])['status'] == 'AMBIGUOUS'


def test_forecast_saving_is_not_an_amount_anchor_and_legacy_ids_are_not_proof():
    rec = recommendation('Рекомендуем позицию 42 вынести на ЭА (Поставка бумаги) на сумму 100,00 тыс. руб. (прогнозная экономия — 46,00 тыс. руб.)')
    assert resolve(rec, [row()])['status'] != 'CONFIRMED'
    rec = recommendation('Рекомендуем поставить бумагу на сумму 46,00 тыс. руб.')
    rec['source_procurement_ids'] = ['42']
    assert resolve(rec, [row()])['status'] != 'CONFIRMED'


def test_group_target_and_recorded_fact_do_not_establish_group_fulfillment():
    text = 'Объединить позиции 42, 43 в единую закупку Поставка бумаги на сумму 46,00 тыс. руб.'
    result = resolve(recommendation(text), [row(actual_date='2026-09-29')])
    assert result['status'] != 'CONFIRMED'
    assert result['fulfillment'] == 'UNKNOWN'


def test_origin_text_hash_or_future_origin_date_cannot_be_accepted():
    rec = recommendation(); rec['origin_evidence'][0]['text_sha256'] = '0' * 64
    assert resolve(rec, [row()])['status'] == 'ORIGIN_UNPROVEN'
    rec = recommendation(); rec['origin_evidence'][0]['document_date'] = '2026-10-01'
    assert resolve(rec, [row()])['status'] == 'ORIGIN_UNPROVEN'


def test_an_unresolved_duplicate_must_not_make_the_other_candidate_unique():
    assert resolve(recommendation(), [row(), replace(row(), row_number=5, procurement_uid=None)])['status'] == 'AMBIGUOUS'


@pytest.mark.parametrize('text', [
    'Вынести на ЭА 42 Поставка бумаги на сумму 46,00 тыс. руб.',
    '42 Поставка бумаги на сумму 46,00 тыс. руб.',
])
def test_explicit_number_after_requested_method_or_at_record_start_is_supported(text):
    assert resolve(recommendation(text), [row()])['status'] == 'CONFIRMED'


@pytest.mark.parametrize('text', [
    '42,00 тыс. руб. Поставка бумаги',
    '42 тыс. руб. Поставка бумаги',
    '42.09.2026 Поставка бумаги на сумму 46,00 тыс. руб.',
])
def test_amounts_and_dates_at_text_start_are_not_position_numbers(text):
    assert resolve(recommendation(text), [row()])['status'] != 'CONFIRMED'


@pytest.mark.parametrize('year', [None, 2026])
def test_future_target_year_cannot_confirm_a_current_year_or_undated_plan(year):
    rec = recommendation(TEXT + ' Перенести из плана 2026 года в план 2027 года.')
    assert resolve(rec, [replace(row(), planned_year=year)])['status'] != 'CONFIRMED'


def test_subject_is_a_full_phrase_not_a_substring_of_another_subject():
    rec = recommendation('Рекомендуем позицию 42 вынести на ЭА (Экраны) на сумму 46,00 тыс. руб.')
    assert resolve(rec, [replace(row(), subject='Краны')])['status'] != 'CONFIRMED'


@pytest.mark.parametrize('text', [
    TEXT.replace('Поставка бумаги)', 'Поставка бумаги и картриджей)'),
    'Поставка бумаги отменена. Рекомендуем позицию 42 (Поставка картриджей) на сумму 46,00 тыс. руб.',
    '42 Поставка бумаги и картриджей на сумму 46,00 тыс. руб.',
])
def test_subject_belongs_in_full_to_the_referenced_position(text):
    assert resolve(recommendation(text), [row()])['status'] != 'CONFIRMED'


def test_complete_subject_with_dash_before_its_amount_is_supported():
    assert resolve(recommendation('Вынести на ЭА 42 Поставка бумаги – 46,00 тыс. руб.'), [row()])['status'] == 'CONFIRMED'


@pytest.mark.parametrize('subject,expected', [
    ('оргтехника ( 2 принтера, 2 ноутбука)', 'CONFIRMED'),
    ('оргтехника (2 принтера, 2 монитора)', 'CURRENT_EVIDENCE_MISSING'),
    ('оргтехника (2 принтера, 2 ноутбука) и картриджи', 'CURRENT_EVIDENCE_MISSING'),
])
def test_v9_literal_reference_tolerates_parenthesis_spacing_only(subject, expected):
    from procurement_engine.raw_pipeline import review_recommendations

    rec = recommendation('Вынести на ЭА 42 оргтехника (2 принтера, 2 ноутбука) – 180,00 тыс. руб.')
    rec['active_in_current_slice'] = True
    rows = [replace(row(), subject=subject, planned_year=2026)]
    kwargs = {'documents': TEST_DOCUMENTS, 'budget_years': {}}
    result = review_recommendations([rec], rows, 'snapshot', '30.09.2026',
        link_contract='verified-original-and-current-plan-v9', **kwargs)
    assert result[0]['current_link']['status'] == expected
    old = review_recommendations([rec], rows, 'snapshot', '30.09.2026',
        link_contract='verified-original-and-current-plan-v8', **kwargs)
    assert old[0]['current_link']['status'] == 'CURRENT_EVIDENCE_MISSING'


@pytest.mark.parametrize('change,expected', [
    ({}, 'CONFIRMED'),
    ({'subject': 'Поставка бумаги и картриджей'}, 'GROUP_EVIDENCE_REQUIRED'),
    ({'procurement_uid': None}, 'GROUP_EVIDENCE_REQUIRED'),
    ({'source_row_no': '44'}, 'GROUP_EVIDENCE_REQUIRED'),
])
def test_v9_group_target_parenthesis_keeps_full_original_members(change, expected):
    from procurement_engine.raw_pipeline import review_recommendations

    rec = recommendation('Вынести на ЭА 42,43 (99) Поставка бумаги -100,00 тыс. руб. (совместный аукцион)')
    rec['active_in_current_slice'] = True
    second = {'source_row_no': '43', 'procurement_id': '43', 'row_number': 5,
              'procurement_uid': 'PUR-other', 'planned_year': 2026} | change
    rows = [replace(row(), planned_year=2026), replace(row(), **second)]
    result = review_recommendations([rec], rows, 'snapshot', '30.09.2026',
        documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v9', budget_years={})
    assert result[0]['current_link']['status'] == expected
    if expected == 'CONFIRMED':
        assert result[0]['current_link']['business_ids'] == ['42', '43']
        assert result[0]['current_link']['fulfillment'] == 'UNKNOWN'
        assert 'relation' not in result[0]['current_link']
        duplicate = replace(rows[1], row_number=6, procurement_uid=None)
        result = review_recommendations([rec], [*rows, duplicate], 'snapshot', '30.09.2026',
            documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v9', budget_years={})
        assert result[0]['current_link']['status'] == 'AMBIGUOUS'


@pytest.mark.parametrize('text', [
    'Рекомендуем позицию 42 (Поставка бумаги) и картриджей на сумму 46,00 тыс. руб.',
    'Рекомендуем позицию 42 (Поставка бумаги). Поставка картриджей на сумму 46,00 тыс. руб.',
])
def test_closed_subject_does_not_borrow_another_subjects_amount(text):
    assert resolve(recommendation(text), [row()])['status'] != 'CONFIRMED'


@pytest.mark.parametrize('amount', ['100,00 тыс. руб.', ', которую ещё нужно уточнить.'])
def test_planned_money_elsewhere_cannot_override_the_positions_own_amount(amount):
    text = ('Рекомендуем позицию 42 (Поставка бумаги) на сумму ' + amount
            + ' Для поставки картриджей плановая сумма 46,00 тыс. руб.')
    assert resolve(recommendation(text), [row()])['status'] != 'CONFIRMED'


def test_composite_business_number_is_not_truncated():
    rec = recommendation(TEXT.replace('позицию 42', 'позицию 42/1'))
    assert resolve(rec, [row()])['status'] != 'CONFIRMED'
    assert resolve(rec, [replace(row(), source_row_no='42/1', procurement_id='42/1')])['status'] == 'CONFIRMED'


def test_quantity_at_text_start_is_not_a_procurement_number():
    rec = recommendation('42 пачки. Поставка бумаги на сумму 46,00 тыс. руб.')
    assert resolve(rec, [row()])['status'] != 'CONFIRMED'


def test_expected_saving_with_intervening_words_is_not_the_procurement_amount():
    rec = recommendation('Рекомендуем позицию 42 Поставка бумаги на сумму 100,00 тыс. руб. Экономия составит 46,00 тыс. руб.')
    assert resolve(rec, [row()])['status'] != 'CONFIRMED'


def test_four_digit_business_number_does_not_become_a_year():
    rec = recommendation(TEXT.replace('позицию 42', 'позицию 2036'))
    assert resolve(rec, [replace(row(), source_row_no='2036', procurement_id='2036')])['status'] == 'CONFIRMED'


def test_letter_suffix_is_not_truncated_to_another_business_number():
    rec = recommendation(TEXT.replace('позицию 42', 'позицию 42а'))
    assert resolve(rec, [row()])['status'] != 'CONFIRMED'
    assert resolve(rec, [replace(row(), source_row_no='42а')])['status'] == 'CONFIRMED'


@pytest.mark.parametrize('period', ['в 2027 году', 'в 2027-м году'])
def test_future_year_inflections_cannot_confirm_the_current_plan(period):
    assert resolve(recommendation(TEXT + ' Закупку провести ' + period), [row()])['status'] != 'CONFIRMED'


@pytest.mark.parametrize('text', [
    'Рекомендуем позицию 42 Поставка бумаги. Плановая сумма 100,00 тыс. руб. Сумма заключенного договора 46,00 тыс. руб.',
    TEXT.replace('46,00', '-46,00'),
])
def test_contract_amount_and_negative_amount_are_not_positive_plan_anchors(text):
    assert resolve(recommendation(text), [row()])['status'] != 'CONFIRMED'


def test_number_sign_after_position_marker_is_supported():
    assert resolve(recommendation(TEXT.replace('позицию 42', 'позицию №42')), [row()])['status'] == 'CONFIRMED'


@pytest.mark.parametrize('amount', ['Оплата 46,00', 'Факт 46,00'])
def test_explicit_plan_amount_has_priority_over_other_amounts(amount):
    text = 'Рекомендуем позицию 42 Поставка бумаги. Плановая сумма 100,00 тыс. руб. ' + amount + ' тыс. руб.'
    assert resolve(recommendation(text), [row()])['status'] != 'CONFIRMED'
    assert resolve(recommendation(text), [replace(row(), plan_mb=100)])['status'] == 'CONFIRMED'


@pytest.mark.parametrize('negative', ['- 46,00', '−46,00'])
def test_negative_amount_with_space_or_unicode_minus_is_not_positive(negative):
    assert resolve(recommendation(TEXT.replace('46,00', negative)), [row()])['status'] != 'CONFIRMED'


def saved_report(grbs='УЭР'):
    from io import BytesIO

    from docx import Document

    doc = Document(); doc.add_paragraph('ОТЧЕТ ПО ЗАКУПКАМ'); doc.add_paragraph('срез на 25.09.2026')
    doc.add_paragraph(grbs)
    table = doc.add_table(rows=2, cols=2); table.cell(1, 1).text = TEXT
    buffer = BytesIO(); doc.save(buffer); content = buffer.getvalue()
    rec = recommendation(); e = rec['origin_evidence'][0]
    e.update(document_sha256=hashlib.sha256(content).hexdigest(), table=1, row=2, cell=2, grbs_heading=grbs)
    return rec, content


def test_origin_verification_reads_the_original_document_bytes_and_cell():
    from procurement_engine.recommendation_links import verify_saved_report_origin

    rec, content = saved_report(); digest = hashlib.sha256(content).hexdigest()
    proof = verify_saved_report_origin(rec, {digest: content})
    assert proof['document_sha256'] == digest
    assert proof['text_sha256'] == hashlib.sha256(TEXT.encode()).hexdigest()
    assert (proof['table'], proof['row'], proof['cell']) == (1, 2, 2)


@pytest.mark.parametrize('mutation', ['bytes', 'coordinate', 'department', 'date', 'text', 'missing'])
def test_wrong_original_provenance_cannot_be_registered(mutation):
    from procurement_engine.recommendation_links import verify_saved_report_origin

    rec, content = saved_report('УО' if mutation == 'department' else 'УЭР')
    digest = hashlib.sha256(content).hexdigest()
    if mutation == 'bytes': content = content + b'changed'
    if mutation == 'coordinate': rec['origin_evidence'][0]['row'] = 1
    if mutation == 'date': rec['origin_evidence'][0]['document_date'] = '2026-09-24'
    if mutation == 'text': rec['recommendation_text'] = 'Другая рекомендация'
    docs = {} if mutation == 'missing' else {digest: content}
    assert verify_saved_report_origin(rec, docs) is None


@pytest.mark.parametrize('heading,grbs', [('УД АЕМР + МКУ «ЕДДС»', 'УД'),
    ('УКСиМП АЕМР + подведомственные учреждения', 'УКСиМП'), ('УО АЕМР + подведомственные учреждения', 'УО')])
def test_saved_report_extended_department_headings_are_recognized(heading, grbs):
    from procurement_engine.recommendation_links import verify_saved_report_origin

    rec, content = saved_report(heading); rec['grbs'] = grbs
    assert verify_saved_report_origin(rec, {hashlib.sha256(content).hexdigest(): content}) is not None


@pytest.mark.parametrize('mutation', ['comparison_date', 'body_department_reference'])
def test_body_references_are_not_report_date_or_department_proof(mutation):
    from io import BytesIO

    from docx import Document
    from procurement_engine.recommendation_links import verify_saved_report_origin

    rec, content = saved_report('УО' if mutation == 'body_department_reference' else 'УЭР')
    doc = Document(BytesIO(content))
    if mutation == 'comparison_date':
        doc.paragraphs[1].text = 'Отчёт на 30.09.2026. Сравнение с отчётом на 25.09.2026.'
    else:
        p = doc.add_paragraph('УЭР просит проверить рекомендации УО.')
        doc.tables[0]._tbl.addprevious(p._p)
        rec['origin_evidence'][0]['grbs_heading'] = p.text
    output = BytesIO(); doc.save(output); content = output.getvalue()
    digest = hashlib.sha256(content).hexdigest(); rec['origin_evidence'][0]['document_sha256'] = digest
    assert verify_saved_report_origin(rec, {digest: content}) is None


def test_ledger_metadata_alone_cannot_establish_a_current_link():
    assert resolve_current_link(recommendation(), [row()], report_date='30.09.2026',
                                snapshot_id='snapshot')['status'] == 'ORIGIN_UNPROVEN'


@pytest.mark.parametrize('digest', [[], {}, None, 42])
def test_malformed_origin_digest_is_an_unproven_origin(digest):
    from procurement_engine.recommendation_links import verify_saved_report_origin

    rec = recommendation(); rec['origin_evidence'][0]['document_sha256'] = digest
    assert verify_saved_report_origin(rec, TEST_DOCUMENTS) is None


def test_production_review_confirms_link_from_original_bytes_without_claiming_execution():
    from procurement_engine.raw_pipeline import review_recommendations

    rec = recommendation()
    rec.update(active_in_current_slice=True, uer_decision_original='Принята', table_no=1, row_no=1)
    result = review_recommendations([rec], [row()], 'snapshot', '30.09.2026', documents=TEST_DOCUMENTS)[0]
    assert result['current_link']['status'] == 'CONFIRMED'
    assert result['semantic_status'] == 'CURRENT_LINK_CONFIRMED'
    assert result['current_procurement_ids'] == ['42']
    assert result['current_procurement_state'] == 'UNKNOWN'
    assert result['dimensions']['execution_status'] == 'UNKNOWN'
    assert result['uer_decision_original'] == 'Принята'
    assert result['semantic_status_ru'] == ''


def test_production_review_keeps_unproven_origin_in_diagnostics_without_blanket_executive_comment():
    from procurement_engine.raw_pipeline import review_recommendations

    rec = recommendation(); rec['active_in_current_slice'] = True
    result = review_recommendations([rec], [row()], 'snapshot', '30.09.2026')[0]
    assert result['current_link']['status'] == 'ORIGIN_UNPROVEN'
    assert result['semantic_status'] == 'REVIEW_REQUIRED'
    assert result['semantic_status_ru'] == ''
    assert result['current_procurement_ids'] == []


def test_same_original_document_is_parsed_once_for_multiple_verifications(monkeypatch):
    from procurement_engine import recommendation_links

    rec = recommendation(TEXT + ' Проверочный уникальный текст.')
    original = recommendation_links.Document; calls = []
    def counted(content):
        calls.append(1)
        return original(content)
    monkeypatch.setattr(recommendation_links, 'Document', counted)
    assert recommendation_links.verify_saved_report_origin(rec, TEST_DOCUMENTS)
    assert recommendation_links.verify_saved_report_origin(rec, TEST_DOCUMENTS)
    assert len(calls) == 1


def test_complete_group_links_each_own_subject_amount_and_uid_without_claiming_merge():
    text = ('Объединить позиции 42, 43 в одну закупку. '
            'Позиция 42 (Поставка бумаги) на сумму 46,00 тыс. руб.; '
            'позиция 43 (Поставка картриджей) на сумму 17,00 тыс. руб.')
    second = replace(row(), source_row_no='43', procurement_id='43', subject='Поставка картриджей',
                     plan_mb=17, procurement_uid='PUR-second', row_number=5)
    result = resolve(recommendation(text), [row(), second])
    assert result['status'] == 'CONFIRMED'
    assert result['business_ids'] == ['42', '43']
    assert result['procurement_uids'] == ['PUR-synthetic', 'PUR-second']
    assert result['fulfillment'] == 'UNKNOWN'
    assert result['origin']['document_sha256']


@pytest.mark.parametrize('change', [{'plan_mb': 63}, {'subject': 'Другой предмет'},
    {'procurement_uid': None}, {'procurement_uid': 'PUR-synthetic'}])
def test_group_requires_every_member_not_combined_money_or_reused_identity(change):
    text = ('Объединить позиции 42, 43 в одну закупку. '
            'Позиция 42 (Поставка бумаги) на сумму 46,00 тыс. руб.; '
            'позиция 43 (Поставка картриджей) на сумму 17,00 тыс. руб.')
    second = replace(row(), source_row_no='43', subject='Поставка картриджей', plan_mb=17,
                     procurement_uid='PUR-second', row_number=5)
    assert resolve(recommendation(text), [row(), replace(second, **change)])['status'] != 'CONFIRMED'


def test_previous_single_only_link_contract_keeps_its_original_replay_result():
    from procurement_engine.raw_pipeline import review_recommendations

    text = ('Объединить позиции 42, 43 в одну закупку. '
            'Позиция 42 (Поставка бумаги) на сумму 46,00 тыс. руб.; '
            'позиция 43 (Поставка картриджей) на сумму 17,00 тыс. руб.')
    rec = recommendation(text); rec['active_in_current_slice'] = True
    second = replace(row(), source_row_no='43', subject='Поставка картриджей', plan_mb=17,
                     procurement_uid='PUR-second', row_number=5)
    old = review_recommendations([rec], [row(), second], 'snapshot', '30.09.2026',
        documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v1')[0]
    assert old['current_link']['status'] == 'GROUP_EVIDENCE_REQUIRED'
    assert old['semantic_status'] == 'REVIEW_REQUIRED'
    new = review_recommendations([rec], [row(), second], 'snapshot', '30.09.2026',
        documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v2')[0]
    assert new['current_link']['status'] == 'CONFIRMED'
    assert new['dimensions']['grouping_status'] == 'UNKNOWN'


def normalized_original(text=TEXT):
    from io import BytesIO

    from docx import Document
    rec, content = saved_report()
    doc = Document(BytesIO(content))
    observed = text.replace('Поставка бумаги', 'Поставка\nбумаги').replace('46,00', '46,00\u00a0') + '  '
    doc.tables[0].cell(1, 1).text = observed
    out = BytesIO(); doc.save(out); content = out.getvalue()
    e = rec['origin_evidence'][0]
    e.update(document_sha256=hashlib.sha256(content).hexdigest(),
        normalization='NFKC_WHITESPACE_V1', observed_text=observed,
        observed_text_sha256=hashlib.sha256(observed.encode()).hexdigest())
    return rec, content


def test_explicit_typographic_origin_keeps_original_bytes_and_cell():
    from procurement_engine.recommendation_links import verify_saved_report_origin
    rec, content = normalized_original()
    proof = verify_saved_report_origin(rec, {hashlib.sha256(content).hexdigest(): content})
    assert proof is not None
    assert proof['text_sha256'] == hashlib.sha256(TEXT.encode()).hexdigest()


@pytest.mark.parametrize('change', ['other-word', 'other-money', 'other-number', 'negated', 'raw-hash', 'raw-cell', 'rule', 'no-rule'])
def test_typographic_equivalence_cannot_change_semantic_source_or_drop_raw_proof(change):
    from procurement_engine.recommendation_links import verify_saved_report_origin
    text = TEXT
    if change == 'other-word': text = text.replace('бумаги', 'картриджей')
    if change == 'other-money': text = text.replace('46,00', '64,00')
    if change == 'other-number': text = text.replace('42', '24')
    if change == 'negated': text = 'Не ' + text
    rec, content = normalized_original(text)
    evidence = rec['origin_evidence'][0]
    if change == 'raw-hash': evidence['observed_text_sha256'] = '0' * 64
    if change == 'raw-cell': evidence['observed_text'] = 'forged'
    if change == 'rule': evidence['normalization'] = 'fuzzy'
    if change == 'no-rule': evidence.pop('normalization')
    assert verify_saved_report_origin(rec, {hashlib.sha256(content).hexdigest(): content}) is None


def test_v4_evaluates_whole_reference_while_frozen_v3_keeps_previous_action_result():
    from procurement_engine.raw_pipeline import review_recommendations
    text = 'Вынести на ЭА 42 Поставка бумаги – 46,00 тыс. руб.'
    rec = recommendation(text); rec['active_in_current_slice'] = True
    current = [replace(row(), planned_year=2026)]
    old = review_recommendations([rec], current, 'snapshot', '30.09.2026', documents=TEST_DOCUMENTS,
                                link_contract='verified-original-and-current-plan-v3')[0]
    new = review_recommendations([rec], current, 'snapshot', '30.09.2026', documents=TEST_DOCUMENTS,
                                link_contract='verified-original-and-current-plan-v4')[0]
    assert old['dimensions']['compliance_status'] == 'UNKNOWN'
    assert new['dimensions']['compliance_status'] == 'IMPLEMENTED'
    assert new['dimensions']['execution_status'] == 'PLANNED'


def test_legacy_exact_origin_with_research_normalization_metadata_is_still_exact():
    from procurement_engine.recommendation_links import verify_saved_report_origin
    rec, content = saved_report()
    rec['origin_evidence'][0]['normalization'] = 'NFKC_casefold_yo_whitespace_v1'
    assert verify_saved_report_origin(rec, {hashlib.sha256(content).hexdigest(): content}) is not None


def resolve_subject_only(rec, rows):
    from procurement_engine.recommendation_links import verify_saved_report_origin

    proof = verify_saved_report_origin(rec, TEST_DOCUMENTS)
    return resolve_current_link(rec, rows, report_date='30.09.2026', snapshot_id='snapshot',
        verified_origin=proof, entity_link_rules=True, exact_subject_fallback=True)


def test_v6_shared_full_subject_links_every_explicit_group_member_and_preserves_v5():
    from procurement_engine.raw_pipeline import review_recommendations

    rec = recommendation('Вынести на ЭА 42,43 Поставка бумаги – 92,00 тыс. руб. (совместный аукцион)')
    rec['active_in_current_slice'] = True
    rows = [replace(row(), planned_year=2026), replace(row(), source_row_no='43',
        procurement_id='43', row_number=5, procurement_uid='PUR-second', planned_year=2026)]
    old = review_recommendations([rec], rows, 'snapshot', '30.09.2026', documents=TEST_DOCUMENTS,
        link_contract='verified-original-and-current-plan-v5')[0]
    current = review_recommendations([rec], rows, 'snapshot', '30.09.2026', documents=TEST_DOCUMENTS,
        link_contract='verified-original-and-current-plan-v6')[0]
    assert old['current_link']['status'] == 'GROUP_EVIDENCE_REQUIRED'
    assert current['current_link']['status'] == 'CONFIRMED'
    assert set(current['current_procurement_uids']) == {'PUR-synthetic', 'PUR-second'}
    assert current['current_link']['fulfillment'] == 'UNKNOWN'


@pytest.mark.parametrize('change', [
    {'subject': 'Поставка бумаги и картриджей'}, {'procurement_uid': None},
    {'planned_year': 2027}, {'source_row_no': '44'}, {'procurement_uid': 'PUR-synthetic'},
])
def test_v6_shared_subject_never_confirms_incomplete_or_mismatched_group(change):
    from procurement_engine.raw_pipeline import review_recommendations

    rec = recommendation('Вынести на ЭА 42,43 Поставка бумаги – 92,00 тыс. руб. (совместный аукцион)')
    rec['active_in_current_slice'] = True
    second = replace(row(), source_row_no='43', procurement_id='43', row_number=5,
        procurement_uid='PUR-second', planned_year=2026)
    result = review_recommendations([rec], [replace(row(), planned_year=2026), replace(second, **change)],
        'snapshot', '30.09.2026', documents=TEST_DOCUMENTS,
        link_contract='verified-original-and-current-plan-v6')[0]
    assert result['current_link']['status'] != 'CONFIRMED'


def subject_only_recommendation(text):
    rec = recommendation(text)
    rec['source_procurement_ids'] = []
    return rec


def test_v5_exact_unique_subject_can_link_original_prose_without_a_position_number():
    text = ('Изменить способ определения поставщика с ЕП на ЭА по мероприятию '
            '«Поставка бумаги» 46,00 тыс. руб.')
    result = resolve_subject_only(subject_only_recommendation(text), [replace(row(), planned_year=2026)])
    assert result['status'] == 'CONFIRMED'
    assert result['business_ids'] == ['42']
    assert result['procurement_uids'] == ['PUR-synthetic']
    assert result['matches'][0]['match_basis'] == 'EXACT_DOCUMENT_SUBJECT_AND_CURRENT_UID'
    assert result['matches'][0]['amount_is_identity_key'] is False


def test_v4_keeps_subject_only_original_unlinked_for_replay_compatibility():
    from procurement_engine.recommendation_links import verify_saved_report_origin

    rec = subject_only_recommendation(
        'Изменить способ определения поставщика с ЕП на ЭА по мероприятию «Поставка бумаги» 46,00 тыс. руб.')
    proof = verify_saved_report_origin(rec, TEST_DOCUMENTS)
    result = resolve_current_link(rec, [replace(row(), planned_year=2026)],
        report_date='30.09.2026', snapshot_id='snapshot', verified_origin=proof, entity_link_rules=True)
    assert result['status'] == 'TEXT_REFERENCE_MISSING'


def test_v5_subject_only_link_does_not_use_historical_price_as_identity():
    text = 'Вынести на ЭА Поставка бумаги – 46,00 тыс. руб.'
    result = resolve_subject_only(subject_only_recommendation(text),
                                  [replace(row(), planned_year=2026, plan_mb=999)])
    assert result['status'] == 'CONFIRMED'


@pytest.mark.parametrize('rows, expected', [
    ([replace(row(), planned_year=2027)], 'TEXT_REFERENCE_MISSING'),
    ([replace(row(), planned_year=2026, procurement_uid=None)], 'CURRENT_EVIDENCE_MISSING'),
    ([replace(row(), planned_year=2026, subject='Поставка бумаги специальной')], 'TEXT_REFERENCE_MISSING'),
])
def test_v5_subject_only_link_requires_same_year_full_subject_and_uid(rows, expected):
    rec = subject_only_recommendation('Вынести на ЭА Поставка бумаги – 46,00 тыс. руб.')
    assert resolve_subject_only(rec, rows)['status'] == expected


def test_v5_subject_only_link_rejects_duplicate_exact_subjects():
    current = replace(row(), planned_year=2026)
    other = replace(current, row_number=5, source_row_no='43', procurement_id='43',
                    procurement_uid='PUR-other')
    rec = subject_only_recommendation('Вынести на ЭА Поставка бумаги – 46,00 тыс. руб.')
    assert resolve_subject_only(rec, [current, other])['status'] == 'AMBIGUOUS'


@pytest.mark.parametrize('text', [
    'Объединить раздробленные процедуры по ЭА «Поставка бумаги» 46,00 тыс. руб. в единую закупку.',
    'Вынести на единый ЭА Поставка бумаги – 46,00 тыс. руб.',
    'Совместная закупка Поставка бумаги – 46,00 тыс. руб.',
])
def test_v5_subject_only_link_never_collapses_group_wording_to_one_current_row(text):
    rec = subject_only_recommendation(text)
    result = resolve_subject_only(rec, [replace(row(), planned_year=2026)])
    assert result['status'] == 'GROUP_EVIDENCE_REQUIRED'


def test_v5_subject_only_original_is_linked_and_action_is_checked_end_to_end():
    from procurement_engine.raw_pipeline import review_recommendations

    text = ('Изменить способ определения поставщика с ЕП на ЭА по мероприятию '
            '«Поставка бумаги» 46,00 тыс. руб.')
    rec = subject_only_recommendation(text)
    rec.update(active_in_current_slice=True, uer_decision_original='Принята', table_no=1, row_no=1)
    current = [replace(row(), planned_year=2026, method='ЭА')]
    result = review_recommendations([rec], current, 'snapshot', '30.09.2026',
        documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v5')[0]
    assert result['current_link']['status'] == 'CONFIRMED'
    assert result['current_link']['matches'][0]['match_basis'] == 'EXACT_DOCUMENT_SUBJECT_AND_CURRENT_UID'
    assert result['dimensions']['compliance_status'] == 'IMPLEMENTED'
    assert result['semantic_status'] == 'IMPLEMENTED'
    assert result['compiled_action']['contract'] == 'original-action-v3'


def test_v5_subject_only_group_original_remains_review_required_end_to_end():
    from procurement_engine.raw_pipeline import review_recommendations

    rec = subject_only_recommendation(
        'Объединить раздробленные процедуры по ЭА «Поставка бумаги» 46,00 тыс. руб. в единую закупку.')
    rec.update(active_in_current_slice=True, table_no=1, row_no=1)
    result = review_recommendations([rec], [replace(row(), planned_year=2026)], 'snapshot', '30.09.2026',
        documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v5')[0]
    assert result['current_link']['status'] == 'GROUP_EVIDENCE_REQUIRED'
    assert result['semantic_status'] == 'REVIEW_REQUIRED'


@pytest.mark.parametrize('method,expected', [('ЭА', 'IMPLEMENTED'), ('ЕП', 'NOT_IMPLEMENTED')])
def test_v7_unclosed_outer_quote_keeps_exact_subject_and_old_replay(method, expected):
    from procurement_engine.raw_pipeline import review_recommendations

    subject = 'Контроль мероприятий «Маршрут «Озёр»'
    text = ('Изменить способ определения поставщика с ЕП на ЭА по мероприятию '
            '«' + subject + ' 400,00 тыс. руб.')
    rec = subject_only_recommendation(text)
    rec.update(active_in_current_slice=True, table_no=1, row_no=1)
    current = [replace(row(), subject=subject, planned_year=2026, method=method)]
    old = review_recommendations([rec], current, 'snapshot', '30.09.2026',
        documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v6')[0]
    assert old['current_link']['status'] == 'CONFIRMED'
    assert old['dimensions']['compliance_status'] == 'UNKNOWN'
    result = review_recommendations([rec], current, 'snapshot', '30.09.2026',
        documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v7')[0]
    assert result['current_link']['status'] == 'CONFIRMED'
    assert result['dimensions']['compliance_status'] == expected
    assert result['compiled_action']['contract'] == 'original-action-v4'
    assert result['compiled_action']['source_text'] == text


def joint_row(**changes):
    base = replace(row(), source_row_no='2209', procurement_id='2209',
        institution='Совместные закупки', subject='Мягкий инвентарь',
        plan_mb=10939.5, planned_year=None, method='ЭА',
        procurement_uid='PUR-joint', row_number=1904)
    return replace(base, **changes)


def test_v8_unique_joint_procurement_proves_group_target_without_reusing_source_numbers():
    from procurement_engine.raw_pipeline import review_recommendations

    text = ('Объединить позиции 1446, 1737, 1774, 1630 мягкий инвентарь '
            'на общую сумму 10 939,50 тыс. руб. в совместную закупку')
    rec = recommendation(text)
    rec.update(active_in_current_slice=True, recommendation_type='MERGE_PROCUREMENTS',
               table_no=1, row_no=1)
    current = [joint_row()]
    old = review_recommendations([rec], current, 'snapshot', '30.09.2026',
        documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v7')[0]
    assert old['current_link']['status'] == 'GROUP_EVIDENCE_REQUIRED'

    result = review_recommendations([rec], current, 'snapshot', '30.09.2026',
        documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v8')[0]
    assert result['current_link']['status'] == 'CONFIRMED'
    assert result['current_link']['relation'] == 'MERGES_INTO'
    assert result['current_link']['required_business_ids'] == ['1446', '1737', '1774', '1630']
    assert result['current_procurement_ids'] == ['2209']
    assert result['current_procurement_uids'] == ['PUR-joint']
    assert result['dimensions']['compliance_status'] == 'IMPLEMENTED'
    assert result['dimensions']['grouping_status'] == 'MERGED'
    assert result['compiled_action']['contract'] == 'original-action-v5'


@pytest.mark.parametrize('change', [
    {'institution': 'МБОУ Школа'},
    {'plan_mb': 10939.49},
    {'method': 'ЕП'},
    {'planned_year': 2027},
    {'procurement_uid': None},
    {'subject': 'Мягкий инвентарь и мебель'},
])
def test_v8_joint_target_rejects_non_primary_or_changed_target(change):
    from procurement_engine.raw_pipeline import review_recommendations

    text = ('Объединить позиции 1446, 1737, 1774, 1630 мягкий инвентарь '
            'на общую сумму 10 939,50 тыс. руб. в совместную закупку')
    rec = recommendation(text)
    rec.update(active_in_current_slice=True, recommendation_type='MERGE_PROCUREMENTS',
               table_no=1, row_no=1)
    result = review_recommendations([rec], [joint_row(**change)], 'snapshot', '30.09.2026',
        documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v8')[0]
    assert result['current_link']['status'] != 'CONFIRMED'



def test_v8_duplicate_exact_joint_targets_remain_ambiguous():
    from procurement_engine.raw_pipeline import review_recommendations

    text = ('Объединить позиции 1446, 1737, 1774, 1630 мягкий инвентарь '
            'на общую сумму 10 939,50 тыс. руб. в совместную закупку')
    rec = recommendation(text)
    rec.update(active_in_current_slice=True, recommendation_type='MERGE_PROCUREMENTS',
               table_no=1, row_no=1)
    first = joint_row()
    second = replace(first, source_row_no='2210', procurement_id='2210',
                     procurement_uid='PUR-joint-2', row_number=1905)
    result = review_recommendations([rec], [first, second], 'snapshot', '30.09.2026',
        documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v8')[0]
    assert result['current_link']['status'] == 'AMBIGUOUS'


@pytest.mark.parametrize('missing', ['H', 'I', 'J'])
def test_v8_joint_target_requires_every_explicit_plan_component(missing):
    from procurement_engine.raw_pipeline import review_recommendations

    text = ('Объединить позиции 1446, 1737, 1774, 1630 мягкий инвентарь '
            'на общую сумму 10 939,50 тыс. руб. в совместную закупку')
    rec = recommendation(text)
    rec.update(active_in_current_slice=True, recommendation_type='MERGE_PROCUREMENTS')
    result = review_recommendations([rec], [joint_row(missing_money_fields=(missing,))],
        'snapshot', '30.09.2026', documents=TEST_DOCUMENTS,
        link_contract='verified-original-and-current-plan-v8')[0]
    assert result['current_link']['status'] != 'CONFIRMED'
    assert result['dimensions']['compliance_status'] == 'UNKNOWN'


def test_v8_unresolved_joint_duplicate_cannot_make_another_target_unique():
    from procurement_engine.raw_pipeline import review_recommendations

    text = ('Объединить позиции 1446, 1737, 1774, 1630 мягкий инвентарь '
            'на общую сумму 10 939,50 тыс. руб. в совместную закупку')
    rec = recommendation(text)
    rec.update(active_in_current_slice=True, recommendation_type='MERGE_PROCUREMENTS')
    duplicate = joint_row(source_row_no='2210', procurement_id='2210',
        procurement_uid=None, row_number=1905)
    result = review_recommendations([rec], [joint_row(), duplicate], 'snapshot',
        '30.09.2026', documents=TEST_DOCUMENTS,
        link_contract='verified-original-and-current-plan-v8')[0]
    assert result['current_link']['status'] == 'AMBIGUOUS'
    assert result['dimensions']['compliance_status'] == 'UNKNOWN'


def test_v8_parenthetical_target_must_match_the_current_joint_business_number():
    from procurement_engine.raw_pipeline import review_recommendations

    text = ('Вынести на ЭА 1205,1309 (2277) оргтехника – 660,00 тыс. руб. '
            '(совместный аукцион)')
    rec = recommendation(text)
    rec.update(active_in_current_slice=True, recommendation_type='MERGE_PROCUREMENTS',
               table_no=1, row_no=1)
    target = replace(joint_row(), source_row_no='9999', procurement_id='9999',
        subject='Оргтехника', plan_mb=660, procurement_uid='PUR-office')
    result = review_recommendations([rec], [target], 'snapshot', '30.09.2026',
        documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v8')[0]
    assert result['current_link']['status'] != 'CONFIRMED'


def test_v8_parenthetical_target_can_prove_the_exact_current_joint_row():
    from procurement_engine.raw_pipeline import review_recommendations

    text = ('Вынести на ЭА 1205,1309 (2277) оргтехника – 660,00 тыс. руб. '
            '(совместный аукцион)')
    rec = recommendation(text)
    rec.update(active_in_current_slice=True, recommendation_type='MERGE_PROCUREMENTS',
               table_no=1, row_no=1)
    target = replace(joint_row(), source_row_no='2277', procurement_id='2277',
        subject='Оргтехника', plan_mb=660, procurement_uid='PUR-office')
    result = review_recommendations([rec], [target], 'snapshot', '30.09.2026',
        documents=TEST_DOCUMENTS, link_contract='verified-original-and-current-plan-v8')[0]
    assert result['current_link']['status'] == 'CONFIRMED'
    assert result['current_procurement_ids'] == ['2277']
    assert result['dimensions']['grouping_status'] == 'MERGED'
