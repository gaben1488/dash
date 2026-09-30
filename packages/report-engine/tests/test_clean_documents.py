import json

from docx import Document
from procurement_engine.docx_renderer import render_main_docx, render_management_docx
from procurement_engine.runtime import run_once
from test_recorded_release import CompleteGoogle
from test_runtime import inputs


def test_clean_word_keeps_business_date_and_hidden_provenance_with_protocol(tmp_path):
    registry, ledger = inputs(tmp_path)
    result = run_once(registry, ledger, tmp_path / 'state', client=CompleteGoogle())
    bundle = tmp_path / 'state/attempts' / result['attempt_id'] / 'bundle'
    model = json.loads((bundle / 'report_model.json').read_text())
    for name in ['main_report.docx', 'management_report.docx']:
        doc = Document(bundle / name)
        text = '\n'.join(p.text for p in doc.paragraphs)
        assert model['snapshot']['report_date'] in text
        assert model['snapshot']['snapshot_id'] not in text
        for token in ['Идентификатор среза', 'версия правил', 'датой Q', 'V–X', 'AD «да»', 'validation layer']:
            assert token not in text
        assert doc.core_properties.identifier == model['snapshot']['snapshot_id']
    protocol = json.loads((bundle / 'diagnostic_protocol.json').read_text())
    assert protocol['snapshot'] == model['snapshot']
    assert protocol['issues'] == model['issues']
    assert protocol['recommendations'] == json.loads((bundle / 'recommendation_review.json').read_text())


def test_future_structured_rows_do_not_get_comment_warning_or_physical_row_label(tmp_path):
    model = {'snapshot': {'report_date': '30.09.2026', 'snapshot_id': 'snapshot', 'report_year': 2026},
             'headline': {'competitive': {'year': {}, 'quarter': {}}, 'single_supplier': {'year': {}, 'quarter': {}}},
             'future_plan': {'target_year': 2027, 'rows': [{'grbs': 'УЭР', 'row_number': 9999,
                 'subject': 'Бумага', 'business_id': '12', 'amount_thousand': 15, 'target_month': 2,
                 'classification': 'STRUCTURED_PLAN_2027', 'review_required': False}]}}
    for render, name in [(render_main_docx, 'main'), (render_management_docx, 'supplement')]:
        path = tmp_path / (name + '.docx')
        render(model, path)
        text = '\n'.join(p.text for p in Document(path).paragraphs)
        assert 'Бумага' in text
        assert '9999' not in text
        assert 'комментари' not in text
        assert 'не установлен' not in text
    model['future_plan']['rows'][0].update(target_month=None, review_required=True,
                                          classification='TARGET_PLAN_2027_UNSTRUCTURED')
    render_main_docx(model, tmp_path / 'uncertain.docx')
    text = '\n'.join(p.text for p in Document(tmp_path / 'uncertain.docx').paragraphs)
    assert 'Бумага' in text and 'не установлен' in text and 'комментари' in text


class BusinessGoogle(CompleteGoogle):
    def grid(self, provider, sheet_id):
        grid = super().grid(provider, sheet_id)
        if provider == 'master-0':
            grid['gridProperties']['rowCount'] = 6
        return grid

    def values(self, provider, title, start, end, columns):
        if provider != 'master-0':
            return super().values(provider, title, start, end, columns)
        rows = [[], [], ['Synthetic header']]
        for number, quarter, subject, amounts, actual in [
            (1, 1, 'Завершённая закупка', [10, 20, 30], True),
            (2, 3, 'Поставка материалов', [7, 8, 9], False),
            (3, 4, 'Закупка следующего квартала', [11, 12, 13], False),
        ]:
            row = [''] * 34
            row[0] = number; row[5] = 'Заказчик'; row[6] = subject
            row[7:11] = [*amounts, sum(amounts)]
            row[11] = 'ЭА'; row[13] = f'01.{(quarter-1)*3+1:02d}.2026'
            row[14:16] = [quarter, 2026]
            row[21:25] = [1, 2, 3, 6] if actual else [0, 0, 0, 0]
            row[25:29] = [0, 0, 0, 0]; row[29] = 'нет'
            if actual:
                row[16] = '02.01.2026'; row[17:19] = [1, 2026]
            else:
                row[31] = 'Извещение размещено 15.09.2026'
            rows.append(row)
        return rows[start-1:end]


def test_complete_business_sections_include_all_quarters_budgets_and_remaining_positions(tmp_path):
    registry, ledger = inputs(tmp_path)
    result = run_once(registry, ledger, tmp_path / 'state', client=BusinessGoogle())
    assert result['status'] == 'VERIFIED'
    bundle = tmp_path / 'state/attempts' / result['attempt_id'] / 'bundle'
    model = json.loads((bundle / 'report_model.json').read_text())
    content = model['report_content']
    assert content['global']['comp']['q1']['exact_decimal']['plan_fb_amount'] == '10.0'
    assert content['by_grbs']['УЭР']['comp']['q3']['plan_count'] == 1
    assert [x['subject'] for x in content['remaining']['УЭР']['comp']['q3']] == ['Поставка материалов']
    doc = Document(bundle / 'main_report.docx')
    text = '\n'.join(p.text for p in doc.paragraphs)
    for quarter in range(1, 5):
        assert f'{quarter} квартал 2026' in text
    assert 'ФБ — 10,00' in text and 'КБ — 20,00' in text and 'МБ — 30,00' in text
    assert 'Поставка материалов' in text and '15.09.2026' in text
    assert 'Закупка следующего квартала' not in text


def test_independent_audit_rejects_budget_and_remaining_population_mutations(tmp_path):
    from copy import deepcopy

    from procurement_engine.independent_audit import audit_model
    from procurement_engine.section_audit import audit_source_sections

    registry, ledger = inputs(tmp_path)
    result = run_once(registry, ledger, tmp_path / 'state', client=BusinessGoogle())
    attempt = tmp_path / 'state/attempts' / result['attempt_id']
    model = json.loads((attempt / 'bundle/report_model.json').read_text())
    capture = json.loads((attempt / 'capture.json').read_text())
    bad = deepcopy(model)
    bad['report_content']['by_grbs']['УЭР']['comp']['q1']['exact_decimal']['plan_fb_amount'] = '11'
    assert not audit_model(capture, bad)['pass']
    bad = deepcopy(model)
    bad['report_content']['remaining']['УЭР']['comp']['q3'] = []
    assert 'remaining_population' in audit_source_sections(capture, bad, ledger=[])


def test_publisher_requires_the_matching_diagnostic_protocol_for_new_reports(tmp_path):
    import pytest
    from procurement_engine.publication_store import PublicationError, PublicationStore

    registry, ledger = inputs(tmp_path)
    result = run_once(registry, ledger, tmp_path / 'state', client=CompleteGoogle())
    bundle = tmp_path / 'state/attempts' / result['attempt_id'] / 'bundle'
    path = bundle / 'diagnostic_protocol.json'
    protocol = json.loads(path.read_text())
    versions = json.loads((bundle / 'snapshot_bundle/bundle.json').read_text())['after']
    path.unlink()
    with pytest.raises(PublicationError):
        PublicationStore(tmp_path / 'other').publish(bundle, read_revisions=lambda: versions)
    protocol['snapshot']['snapshot_id'] = 'another-snapshot'
    path.write_text(json.dumps(protocol))
    with pytest.raises(PublicationError, match='DIAGNOSTIC_PROTOCOL_MISMATCH'):
        PublicationStore(tmp_path / 'other').publish(bundle, read_revisions=lambda: versions)


def test_clean_remaining_comment_excludes_monitoring_internals_and_placeholder_markers():
    from procurement_engine.models import ProcurementRow
    from procurement_engine.report_content import remaining_position

    row = ProcurementRow(snapshot_id='s', procurement_id='42', grbs='УЭР', subject='Бумага', method='ЭА',
        grbs_comment='Извещение размещено 15.09.2026', deviation_reason='X',
        monitoring_note='[сверка кодов] Глазами не проверено')
    comment = remaining_position(row)['comment']
    assert comment == 'Извещение размещено 15.09.2026'
    assert row.monitoring_note == '[сверка кодов] Глазами не проверено'


def test_quality_warning_names_the_procurement_even_when_business_number_is_empty(tmp_path):
    from dataclasses import asdict

    from procurement_engine.adapters import normalize_master_values
    from procurement_engine.qa import validate_master_values

    raw = [''] * 34; raw[6] = 'Поставка мебели без номера'; raw[21] = 10
    rows = normalize_master_values([raw], snapshot_id='s', expected_grbs='УЭР', source_id='book',
                                   sheet_name='ВСЕ', first_sheet_row=4, data_start_row=0)
    issues = [x.as_dict() for x in validate_master_values([raw], grbs='УЭР', source_id='book', sheet_name='ВСЕ')]
    warning = next(x for x in issues if x['code'] == 'MONETARY_FACT_WITHOUT_COMPLETION_DATE')
    assert warning['context']['row_key'] == rows[0].physical_row_key
    model = {'snapshot': {'report_date': '30.09.2026', 'snapshot_id': 's'}, 'issues': [warning],
             'details': [{**asdict(rows[0]), 'physical_row_key': rows[0].physical_row_key}],
             'headline': {'competitive': {'year': {}, 'quarter': {}}, 'single_supplier': {'year': {}, 'quarter': {}}}}
    render_main_docx(model, tmp_path / 'warning.docx')
    text = '\n'.join(p.text for p in Document(tmp_path / 'warning.docx').paragraphs)
    assert 'Поставка мебели без номера' in text
    assert 'закупка № не указан' not in text


def test_native_report_outline_keeps_control_in_competitive_section_and_notices_at_end(tmp_path):
    registry, ledger = inputs(tmp_path)
    result = run_once(registry, ledger, tmp_path / 'state', client=BusinessGoogle())
    bundle = tmp_path / 'state/attempts' / result['attempt_id'] / 'bundle'
    model = json.loads((bundle / 'report_model.json').read_text())
    model['procedures'] = [{'procedure_code': 'ЭА-test', 'stage': 'Объявлена', 'subject': 'Бумага'}]
    model['issues'] = [{'severity': 'WARN', 'code': 'MONETARY_FACT_WITHOUT_COMPLETION_DATE',
                        'context': {'grbs': 'УЭР', 'procurement_id': '9'}}]
    for render, name in [(render_main_docx, 'main'), (render_management_docx, 'supplement')]:
        path = tmp_path / (name + '.docx')
        render(model, path)
        text = '\n'.join(p.text for p in Document(path).paragraphs)
        assert text.index('ПО КОНКУРЕНТНЫМ ЗАКУПКАМ:') < text.index('В РАБОТЕ') < text.index('ЕДИНСТВЕННЫЙ ПОСТАВЩИК:')
        assert text.index('Сведения, требующие уточнения') > text.index('ЕДИНСТВЕННЫЙ ПОСТАВЩИК:')
        if name == 'main':
            assert 'план на год — 3, на квартал — 1, факт — 0' in text
        else:
            assert text.index('ОСТАТОК ТЕКУЩЕГО КВАРТАЛА — КОНКУРЕНТНЫЕ') < text.index('В РАБОТЕ')
