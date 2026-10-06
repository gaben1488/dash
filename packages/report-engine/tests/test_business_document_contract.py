"""Regressions for lost source explanations and falsely verified document contents.
All source values are synthetic. This module never reads production or Google.
"""
import hashlib
import json
from copy import deepcopy
from dataclasses import replace

import pytest
from docx import Document
from procurement_engine.diagnostics import project_diagnostics
from procurement_engine.docx_renderer import render_main_docx, render_management_docx
from procurement_engine.google_adapter import capture_google
from procurement_engine.identity_store import IdentityStore
from procurement_engine.projections import project_dashboard
from procurement_engine.publication_store import PublicationError, PublicationStore
from procurement_engine.raw_pipeline import (
    build_from_capture,
    dump,
    review_recommendations,
)
from procurement_engine.release_gates import validate_recorded_state_model
from test_forensic_qa import rawrow
from test_recorded_release import CompleteGoogle
from test_runtime import inputs


def text_of(path):
    doc = Document(path)
    return '\n'.join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for r in t.rows for c in r.cells])


def build_case(root, *, raw=None, ledger=None, history_package=None):
    root.mkdir(exist_ok=True)
    regpath, _ = inputs(root)
    registry = json.loads(regpath.read_text())
    capture = capture_google(registry, CompleteGoogle())
    capture.update(report_date='30.09.2026', report_year=2026,
                   captured_at='2026-09-30T00:00:00+00:00')
    if raw is None:
        raw = rawrow('42', subject='Поставка бумаги', fact_date='', plan=(0, 0, 46))
        raw[29] = 'нет'
    source = capture['sources'][0]
    source['values'].append(raw)
    source['rows'] = len(source['values'])
    source['formula_evidence']['rows'] = source['rows']
    if history_package is not None:
        capture['recommendation_history_evidence'] = {'metadata': {'version': 'synthetic-v1'}, 'package': history_package}
    model = build_from_capture(capture, registry, ledger or [], root / 'bundle',
                               identity_store=IdentityStore(root / 'identity.sqlite'))
    return model, root / 'bundle', capture


def publish(root):
    after = json.loads((root / 'snapshot_bundle/bundle.json').read_text())['after']
    return PublicationStore(root.parent / 'published').publish(root, read_revisions=lambda: after)


def rewrite(root, model):
    dump(model, root / 'report_model.json')
    dump(project_dashboard(model), root / 'dashboard.json')
    dump(project_diagnostics(model), root / 'diagnostic_protocol.json')
    render_main_docx(model, root / 'main_report.docx')
    render_management_docx(model, root / 'management_report.docx')


@pytest.mark.parametrize('column,field', [(12, 'single_supplier_reason'), (20, 'deviation_reason'),
    (30, 'necessity_reason'), (31, 'grbs_comment'), (32, 'uer_comment'), (33, 'monitoring_note')])
def test_each_explanation_reaches_both_documents_without_completion(tmp_path, column, field):
    row = rawrow('42', subject='Поставка бумаги', fact_date='', plan=(0, 0, 46))
    row[29] = 'нет'
    marker = 'Заказчик сообщил: поставка необходима для работы учреждения.'
    row[column] = marker
    model, root, _ = build_case(tmp_path, raw=row)
    assert model['details'][0][field] == marker
    assert marker in text_of(root / 'main_report.docx')
    assert marker in text_of(root / 'management_report.docx')
    assert model['headline']['single_supplier']['year']['fact_count'] == 0
    assert model['headline']['single_supplier']['year']['remain_count'] == 1
    assert publish(root)['status'] == 'VERIFIED'


def test_missing_trace_inventory_blocks_even_when_artifacts_are_regenerated(tmp_path):
    model, root, _ = build_case(tmp_path)
    model['trace_records'] = []
    assert 'TRACE_CATALOG_MISMATCH' in {i.code for i in validate_recorded_state_model(model, ledger=[])}
    rewrite(root, model)
    with pytest.raises((PublicationError, ValueError)):
        publish(root)


@pytest.mark.parametrize('mutation', ['append', 'remove', 'number', 'header', 'footer', 'hidden'])
def test_word_changes_are_rejected_even_with_an_updated_artifact_hash(tmp_path, mutation):
    _, root, _ = build_case(tmp_path)
    path = root / 'main_report.docx'
    doc = Document(path)
    invented = 'За отчётный период заключено 999999 договоров.'
    if mutation == 'append':
        doc.add_paragraph(invented)
    elif mutation == 'remove':
        doc.paragraphs[3]._element.getparent().remove(doc.paragraphs[3]._element)
    elif mutation == 'number':
        target = next(p for p in doc.paragraphs if '46,00' in p.text)
        target.text = target.text.replace('46,00', '47,00')
    elif mutation in {'header', 'footer'}:
        getattr(doc.sections[0], mutation).paragraphs[0].text = invented
    else:
        doc.paragraphs[3].runs[0].font.hidden = True
    doc.save(path)
    manifest_path = path.with_suffix('.docx.manifest.json')
    meta = json.loads(manifest_path.read_text())
    meta['artifact_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
    dump(meta, manifest_path)
    with pytest.raises(PublicationError, match='DOCUMENT_CONTENT'):
        publish(root)


def test_current_explanation_does_not_overwrite_fact_but_changes_business_finding():
    from test_recommendation_links import TEST_DOCUMENTS, recommendation, row
    rec = recommendation()
    rec.update(active_in_current_slice=True, recommendation_type='CHANGE_METHOD_EA',
               table_no=1, row_no=1)
    current = replace(row(), grbs_comment='Закупка отложена: финансирование отсутствует.')
    found = review_recommendations([rec], [current], 'snapshot', '30.09.2026',
                                  documents=TEST_DOCUMENTS, context_contract='source-context-v1')[0]
    assert found['dimensions']['compliance_status'] == 'IMPLEMENTED'
    assert found['dimensions']['execution_status'] == 'PLANNED'
    assert current.grbs_comment in found['business_finding']
    assert 'сообщил' in found['business_finding'].lower() or 'комментарий' in found['business_finding'].lower()


def test_unknown_action_produces_no_engine_instruction_in_clean_word(tmp_path):
    ledger = [{'recommendation_id': 'REC-TEST', 'grbs': 'УЭР', 'section': 'ep',
               'table_no': 1, 'row_no': 1, 'recommendation_text': 'Рассмотреть организацию общей закупки.',
               'source_procurement_ids': ['42'], 'active_in_current_slice': True,
               'uer_decision_original': 'Предложение рассмотрено.'}]
    model, root, _ = build_case(tmp_path, ledger=ledger)
    assert model['recommendation_records'][0]['dimensions']['compliance_status'] == 'UNKNOWN'
    for name in ('main_report.docx', 'management_report.docx'):
        text = text_of(root / name).casefold()
        for phrase in ('требуется подтверждение связи', 'для проверки действия нужна',
                       'исполнение действия не установлено', 'identity_review_required', 'review_required'):
            assert phrase not in text
    assert 'Предложение рассмотрено.' in text_of(root / 'main_report.docx')


def test_lost_context_field_cannot_be_validated_from_its_own_projection(tmp_path):
    from procurement_engine.section_audit import audit_source_sections
    row = rawrow('42', subject='Поставка бумаги', fact_date='', plan=(0, 0, 46))
    row[12] = 'Услуги единственного оператора.'
    model, _, capture = build_case(tmp_path, raw=row)
    bad = deepcopy(model)
    bad['details'][0]['single_supplier_reason'] = ''
    assert 'source_context' in audit_source_sections(capture, bad, ledger=[])


@pytest.mark.parametrize('column', [12, 30, 31, 33])
def test_comment_only_update_keeps_identity_and_old_publication(tmp_path, column):
    from datetime import datetime
    from zoneinfo import ZoneInfo

    from procurement_engine.runtime import run_once
    now = datetime.now(ZoneInfo('Asia/Kamchatka'))
    registry_path, ledger_path = inputs(tmp_path)
    registry = json.loads(registry_path.read_text())
    provider = registry['sources'][0]['provider_id']

    class Data(CompleteGoogle):
        def __init__(self, comment): self.comment = comment
        def revision(self, book): return self.comment
        def grid(self, book, sheet):
            result = super().grid(book, sheet)
            if book == provider: result['gridProperties']['rowCount'] = 4
            return result
        def values(self, book, title, start, end, columns):
            result = [[], [], ['Synthetic header']]
            if book == provider:
                row = rawrow('42', subject='Поставка бумаги', fact_date='', plan=(0, 0, 46))
                row[13:16] = [f'01.{now.month:02d}.{now.year}', (now.month - 1) // 3 + 1, now.year]
                row[29] = 'нет'; row[column] = self.comment
                result.append(row)
            return result[start - 1:end]

    state = tmp_path / 'state'
    first = run_once(registry_path, ledger_path, state, client=Data('Исходное пояснение заказчика.'))
    assert first['status'] == 'VERIFIED'
    first_root = state / 'published/releases' / first['publication']['release_id']
    old_bytes = (first_root / 'main_report.docx').read_bytes()
    old = json.loads((first_root / 'report_model.json').read_text())
    new_text = 'Заказчик сообщил: закупка отложена, финансирование отсутствует.'
    second = run_once(registry_path, ledger_path, state, client=Data(new_text))
    assert second['status'] == 'VERIFIED'
    new_root = state / 'published/releases' / second['publication']['release_id']
    new = json.loads((new_root / 'report_model.json').read_text())
    assert old['details'][0]['procurement_uid'] == new['details'][0]['procurement_uid']
    assert old['snapshot']['snapshot_id'] != new['snapshot']['snapshot_id']
    assert new['headline']['single_supplier']['year']['fact_count'] == 0
    assert new_text in text_of(new_root / 'main_report.docx')
    assert new_text in text_of(new_root / 'management_report.docx')
    assert (first_root / 'main_report.docx').read_bytes() == old_bytes


@pytest.mark.parametrize('text,visible', [
    ('Требуется подтверждение связи', False), ('[СВЕРКА КОДОВ] глазами не проверено', False),
    ('Необходимо подтвердить сроки поставки у исполнителя.', True),
    ('Отсутствие финансирования не подтверждено, закупка продолжается.', True),
])
def test_technical_cell_marker_is_private_but_business_uncertainty_is_not_erased(tmp_path, text, visible):
    row = rawrow('42', subject='Поставка бумаги', fact_date='', plan=(0, 0, 46))
    row[29] = 'нет'; row[33] = text
    model, root, _ = build_case(tmp_path, raw=row)
    assert model['details'][0]['monitoring_note'] == text
    assert text in (root / 'diagnostic_protocol.json').read_text()
    for name in ('main_report.docx', 'management_report.docx'):
        assert (text in text_of(root / name)) is visible
    assert publish(root)['status'] == 'VERIFIED'


def test_proven_business_finding_and_each_table_cell_survive_publication(tmp_path):
    import base64

    from test_recommendation_links import TEST_DOCUMENTS, recommendation
    rec = recommendation()
    rec.update(active_in_current_slice=True, recommendation_type='CHANGE_METHOD_EA', table_no=1, row_no=1,
               grbs_response_original='Согласны.', uer_decision_original='Принято.')
    digest = rec['origin_evidence'][0]['document_sha256']
    package = {'format': 'aemr-report-recommendation-history-v1',
               'documents': {digest: base64.b64encode(TEST_DOCUMENTS[digest]).decode()}, 'records': [rec]}
    row = rawrow('42', subject='Поставка бумаги', fact_date='', plan=(0, 0, 46))
    row[11] = 'ЭА'; row[29] = 'нет'; row[31] = 'Закупка отложена: финансирование отсутствует.'
    model, root, _ = build_case(tmp_path, raw=row, ledger=[rec], history_package=package)
    finding = model['recommendation_records'][0]['business_finding']
    assert row[31] in finding
    assert finding in Document(root / 'main_report.docx').tables[0].cell(1, 3).text
    table = next(b for b in model['document_plans']['main']['blocks'] if b['kind'] == 'table')
    assert len(table['cells']) == sum(len(row) for row in table['rows'])
    assert all(cell['model_paths'] for cell in table['cells'])
    assert publish(root)['status'] == 'VERIFIED'
    saved = json.loads((root / 'report_model.json').read_text())
    from procurement_engine.document_content import validate_document_plans
    assert validate_document_plans(saved)
    # An altered recommendation cell also fails even when ZIP hashes are updated.
    path = root / 'main_report.docx'; doc = Document(path)
    doc.tables[0].cell(1, 3).text = 'Рекомендация исполнена и договор оплачен.'
    doc.save(path)
    meta_path = path.with_suffix('.docx.manifest.json'); meta = json.loads(meta_path.read_text())
    meta['artifact_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest(); dump(meta, meta_path)
    with pytest.raises(PublicationError, match='DOCUMENT_CONTENT'):
        publish(root)


def test_forged_narrative_is_rejected_even_after_coherent_model_plan_and_doc_rebuild(tmp_path):
    from procurement_engine.document_content import planned_documents
    from procurement_engine.traceability import complete_trace_catalog
    model, root, _ = build_case(tmp_path)
    fake = deepcopy(model['narratives']['GENERIC_TEMPLATE'][0])
    fake.update(block_id='nar.fake', stage='explain', text='В отчётном периоде заключено 999999 договоров.',
                semantic_key='invented-contracts')
    fake['trace']['report_block_id'] = 'nar.fake'
    model['narratives']['GENERIC_TEMPLATE'].append(fake)
    model['trace_records'] = complete_trace_catalog(model)
    model['document_plans'] = planned_documents(model)
    rewrite(root, model)
    with pytest.raises(PublicationError, match='DOMAIN_RELEASE_CONTRACT_FAILED'):
        publish(root)


def test_numeric_trace_cannot_name_a_noncontributing_source_row(tmp_path):
    from procurement_engine.independent_audit import audit_model
    model, _, capture = build_case(tmp_path)
    assert audit_model(capture, model)['pass']
    model['metric_contributors']['headline.single_supplier.year.plan_amount'] = ['invented-row']
    assert not audit_model(capture, model)['pass']


def test_serial_dates_with_comments_are_not_mistaken_for_a_context_mismatch(tmp_path):
    from datetime import date
    row = rawrow('42', subject='Поставка бумаги', plan=(0, 0, 46), fact=(0, 0, 45))
    row[13] = (date(2026, 9, 25) - date(1899, 12, 30)).days
    row[16] = (date(2026, 9, 26) - date(1899, 12, 30)).days
    row[29] = 'нет'; row[12] = 'Пояснение заказчика о способе закупки.'
    model, root, _ = build_case(tmp_path, raw=row)
    assert model['source_context'][0]['planned_date'] == '2026-09-25'
    assert model['source_context'][0]['actual_date'] == '2026-09-26'
    assert model['independent_audit']['pass']
    assert publish(root)['status'] == 'VERIFIED'


@pytest.mark.parametrize('field,changed', [('method', 'ЭА'), ('planned_year', 2027), ('source_row_no', '43')])
def test_context_period_method_and_label_are_checked_against_raw_cells(tmp_path, field, changed):
    from procurement_engine.section_audit import audit_context
    row = rawrow('42', subject='Поставка бумаги', fact_date='', plan=(0, 0, 46))
    row[29] = 'нет'; row[12] = 'Содержательное пояснение способа.'
    model, _, capture = build_case(tmp_path, raw=row)
    assert audit_context(capture, model)
    model['details'][0][field] = changed
    key = 'business_id' if field == 'source_row_no' else field
    model['source_context'][0][key] = changed
    if field == 'planned_year':
        model['source_context'][0]['in_report_year'] = False
    assert not audit_context(capture, model)


def test_a_footer_with_a_nonstandard_package_path_cannot_bypass_content_validation(tmp_path):
    import zipfile

    from lxml import etree
    _, root, _ = build_case(tmp_path)
    path = root / 'main_report.docx'
    doc = Document(path)
    doc.sections[0].footer.paragraphs[0].text = 'Подписано 999999 договоров.'
    doc.save(path)
    with zipfile.ZipFile(path) as archive:
        parts = {info.filename: archive.read(info) for info in archive.infolist()}
    parts['custom/notes.xml'] = parts.pop('word/footer1.xml')
    relationships = etree.fromstring(parts['word/_rels/document.xml.rels'])
    for relation in relationships:
        if relation.get('Target') == 'footer1.xml':
            relation.set('Target', '../custom/notes.xml')
    parts['word/_rels/document.xml.rels'] = etree.tostring(relationships)
    parts['[Content_Types].xml'] = parts['[Content_Types].xml'].replace(b'/word/footer1.xml', b'/custom/notes.xml')
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    meta_path = path.with_suffix('.docx.manifest.json'); meta = json.loads(meta_path.read_text())
    meta['artifact_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest(); dump(meta, meta_path)
    with pytest.raises(PublicationError, match='DOCUMENT_CONTENT'):
        publish(root)


def test_package_relationship_cannot_redirect_word_to_a_different_document(tmp_path):
    import zipfile

    from lxml import etree
    _, root, _ = build_case(tmp_path)
    path = root / 'main_report.docx'
    with zipfile.ZipFile(path) as archive:
        parts = {info.filename: archive.read(info) for info in archive.infolist()}
    relations = etree.fromstring(parts['_rels/.rels'])
    for relation in relations:
        if relation.get('Type', '').endswith('/officeDocument'):
            relation.set('Target', 'custom/empty.xml')
    parts['_rels/.rels'] = etree.tostring(relations)
    parts['custom/empty.xml'] = b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body/></w:document>'
    content = etree.fromstring(parts['[Content_Types].xml'])
    etree.SubElement(content, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
        PartName='/custom/empty.xml', ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml')
    parts['[Content_Types].xml'] = etree.tostring(content)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    meta_path = path.with_suffix('.docx.manifest.json'); meta = json.loads(meta_path.read_text())
    meta['artifact_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest(); dump(meta, meta_path)
    with pytest.raises(PublicationError, match='DOCUMENT_CONTENT'):
        publish(root)
