import base64
import copy
import hashlib
from io import BytesIO

import pytest
from docx import Document
from procurement_engine.recommendation_history import enroll_history_package


def fixture():
    text = 'Рекомендуем позицию 42 (Поставка бумаги) на сумму 46,00 тыс. руб.'
    doc = Document(); doc.add_paragraph('ОТЧЕТ'); doc.add_paragraph('срез на 25.09.2026')
    doc.add_paragraph('УЭР'); doc.add_table(rows=2, cols=2).cell(1, 1).text = text
    buffer = BytesIO(); doc.save(buffer); content = buffer.getvalue()
    digest = hashlib.sha256(content).hexdigest()
    ledger = [{'recommendation_id': 'synthetic', 'grbs': 'УЭР', 'recommendation_text': text,
               'grbs_response_original': 'Согласны', 'uer_decision_original': 'Принята',
               'semantic_status': 'UNKNOWN'}]
    package = {'format': 'aemr-report-recommendation-history-v1',
               'documents': {digest: base64.b64encode(content).decode()},
               'records': [{'recommendation_id': 'synthetic', 'grbs': 'УЭР', 'recommendation_text': text,
                   'origin_evidence': [{'kind': 'SAVED_REPORT_RECOMMENDATION_TEXT',
                       'document_sha256': digest, 'document_date': '2026-09-25',
                       'table': 1, 'row': 2, 'cell': 2, 'grbs_heading': 'УЭР',
                       'text': text, 'text_sha256': hashlib.sha256(text.encode()).hexdigest()}]}]}
    return ledger, package, content


def test_enrollment_verifies_original_bytes_and_preserves_unchanged_history():
    ledger, package, content = fixture(); original = copy.deepcopy(ledger)
    enriched, documents = enroll_history_package(package, ledger)
    assert ledger == original
    assert list(documents.values()) == [content]
    assert enriched[0]['origin_evidence'] == package['records'][0]['origin_evidence']
    assert {k: enriched[0][k] for k in original[0]} == original[0]


@pytest.mark.parametrize('mutation', ['format', 'base64', 'hash', 'unknown_id', 'duplicate_id',
                                     'text', 'department', 'cell', 'date'])
def test_unverified_or_changed_history_package_is_rejected(mutation):
    ledger, package, _ = fixture()
    key = next(iter(package['documents']))
    if mutation == 'format': package['format'] = 'other'
    if mutation == 'base64': package['documents'][key] = 'not base64'
    if mutation == 'hash': package['documents'][key] = base64.b64encode(b'other').decode()
    if mutation == 'unknown_id': package['records'][0]['recommendation_id'] = 'other'
    if mutation == 'duplicate_id': package['records'].append(copy.deepcopy(package['records'][0]))
    if mutation == 'text': package['records'][0]['recommendation_text'] = 'Other'
    if mutation == 'department': package['records'][0]['grbs'] = 'УО'
    if mutation == 'cell': package['records'][0]['origin_evidence'][0]['cell'] = 1
    if mutation == 'date': package['records'][0]['origin_evidence'][0]['document_date'] = '2026-09-24'
    with pytest.raises(ValueError, match='HISTORY_'):
        enroll_history_package(package, ledger)


def test_package_cannot_inject_a_current_execution_status():
    ledger, package, _ = fixture()
    package['records'][0].update(semantic_status='IMPLEMENTED', current_procurement_ids=['42'])
    enriched, _ = enroll_history_package(package, ledger)
    assert enriched[0]['semantic_status'] == 'UNKNOWN'
    assert 'current_procurement_ids' not in enriched[0]


def test_repeated_enrollment_preserves_evidence_without_duplicate_history():
    ledger, package, _ = fixture()
    once, _ = enroll_history_package(package, ledger)
    twice, _ = enroll_history_package(package, once)
    assert twice == once


def test_a_valid_origin_cannot_mask_an_unverified_second_occurrence():
    ledger, package, _ = fixture()
    bad = copy.deepcopy(package['records'][0]['origin_evidence'][0]); bad['cell'] = 1
    package['records'][0]['origin_evidence'].append(bad)
    with pytest.raises(ValueError, match='HISTORY_ORIGIN_UNPROVEN'):
        enroll_history_package(package, ledger)


def test_unreferenced_document_is_not_silently_enrolled_as_report_evidence():
    ledger, package, _ = fixture(); other = b'unreferenced private content'
    package['documents'][hashlib.sha256(other).hexdigest()] = base64.b64encode(other).decode()
    with pytest.raises(ValueError, match='HISTORY_DOCUMENT_SET_UNREGISTERED'):
        enroll_history_package(package, ledger)


class Drive:
    def __init__(self, body, *, count=1, changed=False):
        self.body = body; self.count = count; self.changed = changed
    def _get(self, url, params=None):
        meta = {'id': 'synthetic-history', 'name': 'aemr-report-recommendation-history-v1.json',
                'mimeType': 'application/json', 'version': '1', 'modifiedTime': '2026-09-25T00:00:00Z'}
        if url.endswith('/files'):
            assert "name = 'aemr-report-recommendation-history-v1.json'" in params['q']
            return {'files': [meta] * self.count}
        if params.get('alt') == 'media': return self.body
        return {**meta, 'version': '2' if self.changed else '1'}


def test_google_history_read_checks_package_and_revision_without_writing_ledger():
    from procurement_engine.recommendation_history import read_google_history

    ledger, package, content = fixture(); original = copy.deepcopy(ledger)
    enriched, documents, revision = read_google_history(Drive(package), ledger)
    assert ledger == original
    assert enriched[0]['origin_evidence']
    assert list(documents.values()) == [content]
    assert revision['version'] == '1'


def test_no_registered_history_file_preserves_ledger_without_fabricated_origin():
    from procurement_engine.recommendation_history import read_google_history

    ledger, package, _ = fixture()
    enriched, documents, revision = read_google_history(Drive(package, count=0), ledger)
    assert enriched == ledger
    assert documents == {}
    assert revision is None


@pytest.mark.parametrize('count,changed', [(2, False), (1, True)])
def test_google_history_refuses_duplicate_or_changing_package(count, changed):
    from procurement_engine.recommendation_history import read_google_history

    ledger, package, _ = fixture()
    with pytest.raises(ValueError, match='HISTORY_'):
        read_google_history(Drive(package, count=count, changed=changed), ledger)
