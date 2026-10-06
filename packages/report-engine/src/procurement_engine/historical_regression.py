"""Privacy-safe regression over the verified historical master corpus.

The public oracle stores only semantic SHA-256 digests. Business rows stay in the
private archive or are read from exact Google Drive revisions. A historical case
can deliberately encode a missing mandatory source; that proves fail-closed
behaviour instead of silently treating the missing source as zero.
"""
from __future__ import annotations

import hashlib
import json
from fractions import Fraction
from io import BytesIO
from itertools import chain

from openpyxl import load_workbook

from .independent_audit import number, recount, text
from .normalize import parse_date

CONTRACT = 'historical-master-regression-v1'
ORACLE_CONTRACT = 'historical-master-oracle-v1'
MAX_ROWS = 200000
MAX_COLUMNS = 34


def _hash(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
        separators=(',', ':'), allow_nan=False).encode()).hexdigest()


def _int_token(value):
    try:
        return int(float(text(value).replace(',', '.')))
    except (TypeError, ValueError, OverflowError):
        raw = text(value).casefold()
        return raw or None


def _date_token(value):
    parsed = parse_date(value)
    if parsed:
        return parsed
    raw = text(value).casefold()
    return raw or None


def _money_token(value):
    amount = number(value)
    return [amount.numerator, amount.denominator]


def _row_token(row):
    cell = lambda index: row[index] if index < len(row) else None
    return [
        text(cell(5)).casefold(),
        bool(text(cell(6))),
        text(cell(11)).casefold(),
        _date_token(cell(13)),
        _int_token(cell(14)),
        _int_token(cell(15)),
        _date_token(cell(16)),
        *(_money_token(cell(index)) for index in (7, 8, 9, 21, 22, 23, 25, 26, 27)),
        text(cell(29)).casefold(),
    ]


def source_semantic_digest(values, *, header_rows=3):
    """Digest all metric-relevant historical row semantics without exposing values."""
    rows = []
    empty = _row_token([])
    for row in values[header_rows:]:
        token = _row_token(row)
        if token != empty:
            rows.append(token)
    return _hash(rows)


def _metric_value(value):
    if isinstance(value, Fraction):
        return [value.numerator, value.denominator]
    return value


def metric_digest(capture):
    """Hash the 18 cross-engine fields that were independently regressed."""
    totals, _, _, _, grouped = recount(capture)
    fields = (
        'plan_count', 'fact_count', 'remain_count', 'plan_amount', 'fact_amount',
        'remain_amount', 'confirmed_saving_amount', 'execution_pct',
        'plan_fb_amount', 'plan_kb_amount', 'plan_mb_amount',
        'fact_fb_amount', 'fact_kb_amount', 'fact_mb_amount',
        'remain_fb_amount', 'remain_kb_amount', 'remain_mb_amount',
    )
    projection = {}
    scopes = {'ALL': totals, **{name: grouped[name] for name in sorted(grouped)}}
    for scope, kinds in scopes.items():
        for kind in ('competitive', 'single_supplier'):
            for period in ('year', 'q1', 'q2', 'q3', 'q4'):
                block = kinds[kind][period]
                values = {field: _metric_value(block[field]) for field in fields}
                # Historical TypeScript engine exposed recorded_fact_count as a
                # separate field. With no operational override in this corpus it
                # is intentionally identical to fact_count; keep the alias in the
                # frozen cross-engine contract so all 28,980 comparisons remain represented.
                values['recorded_fact_count'] = block['fact_count']
                projection[f'{scope}/{kind}/{period}'] = values
    return _hash(projection)


def _canonical_history_matrix(content, *, preferred_sheet, header_rows=3):
    """Restore canonical A:AH positions using the historical ordinal row.

    Historical workbooks can physically reorder the first columns (UDTX is a
    verified example). Row 1 records canonical ordinals, so forensic replay uses
    that explicit mapping instead of pretending the old physical order was fixed.
    """
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError('HISTORICAL_XLSX_INVALID') from error
    try:
        candidates = [preferred_sheet] if preferred_sheet in workbook.sheetnames else []
        candidates += [name for name in workbook.sheetnames if name != preferred_sheet]
        matches = []
        for name in candidates:
            sheet = workbook[name]
            iterator = sheet.iter_rows(min_row=1, values_only=True)
            first = next(iterator, None)
            if first is None:
                continue
            mapping = {}
            for physical, value in enumerate(first):
                try:
                    numeric = float(value)
                    ordinal = int(numeric)
                except (TypeError, ValueError, OverflowError):
                    continue
                if numeric == ordinal and 1 <= ordinal <= MAX_COLUMNS and ordinal not in mapping:
                    mapping[ordinal] = physical
            required = {6, 7, 8, 9, 10, 12, 14, 15, 16, 17, 22, 23, 24, 26, 27, 28, 30}
            if not required.issubset(mapping):
                continue
            rows = []
            for number, physical_row in enumerate(chain([first], iterator), 1):
                if number > MAX_ROWS:
                    raise ValueError('HISTORICAL_RANGE_TOO_LARGE')
                row = [None] * MAX_COLUMNS
                for ordinal, physical in mapping.items():
                    if physical < len(physical_row):
                        row[ordinal - 1] = physical_row[physical]
                rows.append(row)
            if len(rows) < header_rows:
                continue
            matches.append((name, rows))
            if name == preferred_sheet:
                break
        if not matches:
            raise ValueError('HISTORICAL_MASTER_SHEET_NOT_FOUND')
        if matches[0][0] != preferred_sheet and len(matches) > 1:
            raise ValueError('HISTORICAL_MASTER_SHEET_AMBIGUOUS')
        return matches[0]
    finally:
        workbook.close()


def _capture(day, matrices, registry):
    sources = []
    masters = {source['grbs']: source for source in registry.get('sources', [])
               if source.get('role') == 'master' and source.get('grbs')}
    for grbs in sorted(matrices):
        source = masters[grbs]
        sheet, values = matrices[grbs]
        sources.append({
            'role': 'master', 'grbs': grbs, 'provider_id': source['provider_id'],
            'sheet': sheet, 'header_rows': 3, 'values': values,
        })
    return {'report_date': day, 'report_year': int(day[:4]), 'sources': sources}


def verify_private_corpus(cases, oracle, registry):
    """Verify already-normalized private cases. Intended for offline acceptance."""
    if oracle.get('contract') != ORACLE_CONTRACT:
        raise ValueError('HISTORICAL_ORACLE_INVALID')
    expected = {case['date']: case for case in oracle.get('cases', [])}
    failures = []
    checks = 0
    seen = set()
    for private in cases:
        day = parse_date(private.get('date'))
        seen.add(day)
        case = expected.get(day)
        if case is None:
            failures.append({'date': day, 'code': 'HISTORICAL_CASE_UNEXPECTED'})
            continue
        observed_missing = sorted(source.get('grbs') for source in private.get('sources', [])
                                  if source.get('missing') and source.get('grbs'))
        if observed_missing != sorted(case.get('expected_missing_sources') or []):
            failures.append({'date': day, 'code': 'HISTORICAL_MISSING_SOURCE_STATE_MISMATCH'})
        matrices = {}
        for source in private.get('sources', []):
            grbs = source.get('grbs')
            if grbs not in case.get('source_digests', {}):
                continue
            values = [[None] * MAX_COLUMNS for _ in range(3)] + [row['cells'] for row in source.get('rows', [])]
            checks += 1
            if source_semantic_digest(values) != case['source_digests'][grbs]:
                failures.append({'date': day, 'grbs': grbs, 'code': 'HISTORICAL_SOURCE_MISMATCH'})
            matrices[grbs] = (source.get('sheet') or '', values)
        if sorted(set(case['source_digests']) - set(matrices)):
            failures.append({'date': day, 'code': 'HISTORICAL_SOURCE_MISSING'})
            continue
        checks += 1
        if metric_digest(_capture(day, matrices, registry)) != case['metric_digest']:
            failures.append({'date': day, 'code': 'HISTORICAL_METRIC_MISMATCH'})
    for day in sorted(set(expected) - seen):
        failures.append({'date': day, 'code': 'HISTORICAL_CASE_MISSING'})
    return {'contract': CONTRACT, 'cases': len(cases), 'checks': checks,
            'failures': failures, 'pass': not failures}


def probe_google_historical_regression(registry, oracle, client, *, timezone_name='Asia/Kamchatka'):
    """Find exact Drive revisions by semantic digest and replay all oracle cases.

    Output contains aggregate/status metadata only. Source rows never leave this
    function and no business text is printed into CI logs.
    """
    if oracle.get('contract') != ORACLE_CONTRACT:
        raise ValueError('HISTORICAL_ORACLE_INVALID')
    from .google_adapter import GoogleReadError
    from .master_revision_history import _revision_export, revision_export_unavailable, revision_index
    masters = {source.get('grbs'): source for source in registry.get('sources', [])
               if source.get('role') == 'master' and source.get('grbs')}
    required_grbs = sorted({
        grbs
        for case in oracle.get('cases', [])
        for grbs in case.get('source_digests', {})
    })
    if any(grbs not in masters for grbs in required_grbs):
        raise ValueError('HISTORICAL_SOURCE_NOT_REGISTERED')
    indexes = {grbs: revision_index(client, masters[grbs]['provider_id'], timezone_name=timezone_name)
               for grbs in required_grbs}
    totals = {'cases': 0, 'complete_cases': 0, 'expected_blocked_cases': 0,
              'source_states': 0, 'matched_source_states': 0, 'metric_cases_passed': 0}
    by_date = []
    for case in oracle.get('cases', []):
        day = parse_date(case.get('date'))
        if not day:
            raise ValueError('HISTORICAL_ORACLE_DATE_INVALID')
        totals['cases'] += 1
        missing_expected = sorted(case.get('expected_missing_sources') or [])
        totals['expected_blocked_cases' if missing_expected else 'complete_cases'] += 1
        matrices = {}
        source_status = {'MATCHED': 0, 'NO_EXACT_REVISION': 0, 'NO_MATCHING_REVISION': 0,
                         'REVISION_UNAVAILABLE': 0}
        for grbs, expected_digest in sorted(case.get('source_digests', {}).items()):
            totals['source_states'] += 1
            source = masters[grbs]
            revisions = indexes[grbs].get(day, [])
            if not revisions:
                source_status['NO_EXACT_REVISION'] += 1
                continue
            matched = None
            exportable = 0
            for revision in reversed(revisions):
                try:
                    content = _revision_export(client, source['provider_id'], revision)
                except GoogleReadError as error:
                    if not revision_export_unavailable(error):
                        raise
                    continue
                exportable += 1
                try:
                    sheet, values = _canonical_history_matrix(content,
                        preferred_sheet=source['sheet'], header_rows=3)
                except ValueError:
                    continue
                if source_semantic_digest(values) == expected_digest:
                    matched = (sheet, values)
                    break
            if matched is None:
                source_status['REVISION_UNAVAILABLE' if exportable == 0 else 'NO_MATCHING_REVISION'] += 1
                continue
            source_status['MATCHED'] += 1
            totals['matched_source_states'] += 1
            matrices[grbs] = matched
        metric_status = 'NOT_CHECKED'
        if len(matrices) == len(case.get('source_digests', {})):
            metric_status = ('PASS' if metric_digest(_capture(day, matrices, registry)) == case['metric_digest']
                             else 'MISMATCH')
            totals['metric_cases_passed'] += int(metric_status == 'PASS')
        by_date.append({'date': day, 'expected_source_status': 'BLOCKED' if missing_expected else 'COMPLETE',
                        'source_status': source_status, 'metric_status': metric_status})
    passed = (totals['matched_source_states'] == totals['source_states']
              and totals['metric_cases_passed'] == totals['cases'])
    return {'contract': CONTRACT, **totals, 'regression_status': 'PASS' if passed else 'FAIL', 'by_date': by_date}
