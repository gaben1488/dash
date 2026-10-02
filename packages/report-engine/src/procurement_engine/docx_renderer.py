from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

BLUE = "95B3D7"
ORANGE = "E36C0A"
GRAY = "7F7F7F"
RED = "C00000"


def _set_font(run, *, size=9, bold=False, italic=False, color: str | None = None) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _write_paragraph(doc: Document, text: str = "", *, size=9, bold=False, italic=False,
               color: str | None = None, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
               keep_with_next=False, space_after=0, first_line_mm: float | None = None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = keep_with_next
    if first_line_mm is not None:
        p.paragraph_format.first_line_indent = Mm(first_line_mm)
    r = p.add_run(str(text))
    _set_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def _set_cell_text(cell, text: str, *, bold=False, size=7, color: str | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text or ""))
    _set_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _setup_main(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(15.0)
    sec.bottom_margin = Mm(15.0)
    sec.left_margin = Mm(15.0)
    sec.right_margin = Mm(15.0)


def _setup_management(doc: Document) -> None:
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(297), Mm(210)
    sec.top_margin = Mm(12.7)
    sec.bottom_margin = Mm(8.0)
    sec.left_margin = Mm(12.7)
    sec.right_margin = Mm(12.7)


def _artifact_manifest(model: dict, *, view: str, narrative_mode: str, output: Path) -> dict:
    payload = json.dumps(model, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return {
        "artifact": output.name,
        "artifact_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "view": view,
        "narrative_mode": narrative_mode,
        "snapshot_id": model["snapshot"].get("snapshot_id"),
        "cutoff_at": model["snapshot"].get("cutoff_at"),
        "rules_version": model["snapshot"].get("rules_version"),
        "renderer_version": model["snapshot"].get("renderer_version"),
        "report_model_version": (model.get("contract") or {}).get("report_model_version"),
        "report_model_sha256": hashlib.sha256(payload).hexdigest(),
        "trace_record_count": len(model.get("trace_records") or []),
    }


def _save_with_manifest(doc: Document, model: dict, output_path: str | Path, *, view: str, narrative_mode: str) -> tuple[Path, Path]:
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.identifier = str(model["snapshot"].get("snapshot_id") or "")
    doc.core_properties.subject = str(model["snapshot"].get("report_date") or "")
    # Business DOCX has no engineering comment property. Provenance remains in
    # the stable identifier, manifest and private diagnostic protocol.
    doc.core_properties.comments = ""
    doc.save(out)
    manifest = _artifact_manifest(model, view=view, narrative_mode=narrative_mode, output=out)
    sidecar = out.with_suffix(out.suffix + ".manifest.json")
    sidecar.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return out, sidecar




def _write_table(doc, rows):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = (Mm(12), Mm(78), Mm(42), Mm(48))
    for column, width in zip(table.columns, widths):
        column.width = width
    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_column.set(qn("w:w"), str(width.twips))
    for index, values in enumerate(rows):
        row = table.rows[0] if index == 0 else table.add_row()
        if index == 0:
            row.height = Mm(9)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            _repeat_table_header(row)
        if sum(len(str(value)) for value in values) <= 1700:
            _prevent_row_split(row)
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = width
            _set_cell_text(cell, str(value), bold=index == 0)
            if index == 0:
                _shade(cell, "D9EAF7")


def _render_plan(model, plan, output_path):
    doc = Document()
    (_setup_main if plan['view'] == 'main' else _setup_management)(doc)
    alignments = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'right': WD_ALIGN_PARAGRAPH.RIGHT,
                  'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}
    for block in plan['blocks']:
        if block['kind'] == 'paragraph':
            style = dict(block['style'])
            if 'align' in style:
                style['align'] = alignments[style['align']]
            _write_paragraph(doc, block['text'], **style)
        elif block['kind'] == 'table':
            _write_table(doc, block['rows'])
        else:
            raise ValueError('DOCUMENT_CONTENT_UNKNOWN_BLOCK')
    return _save_with_manifest(doc, model, output_path, view=plan['view'], narrative_mode=plan['narrative_mode'])


def render_main_docx(report_model, output_path, *, narrative_mode='GENERIC_TEMPLATE'):
    from .document_plan import build_main_plan
    return _render_plan(report_model, build_main_plan(report_model, narrative_mode=narrative_mode), output_path)


def render_management_docx(report_model, output_path, *, narrative_mode='SMART_NARRATIVE'):
    from .document_plan import build_management_plan
    return _render_plan(report_model, build_management_plan(report_model, narrative_mode=narrative_mode), output_path)
